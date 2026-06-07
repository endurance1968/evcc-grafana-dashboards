#!/usr/bin/env python3
"""
Script: import-sma-pv-energy.py
Purpose: Import exported SMA PV daily energy values into VictoriaMetrics as EVCC-compatible per-title daily rollups.
Version: 2026.06.07.5
Last modified: 2026-06-07
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import re
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

SCRIPT_VERSION = "2026.06.07.5"
SCRIPT_LAST_MODIFIED = "2026-06-07"
DEFAULT_METRIC = "evcc_pv_energy_by_title_daily_wh"


@dataclass(frozen=True)
class SmaEnergyRow:
    day: dt.date
    sma_name: str
    title: str
    energy_wh: float


def log(message: str) -> None:
    print(message, flush=True)


def normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def fold_ascii(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return normalized.encode("ascii", "ignore").decode("ascii")


def normalize_sma_name(value: str) -> str:
    folded = fold_ascii(value).strip().lower()
    return re.sub(r"[^a-z0-9.-]+", "_", folded).strip("_")


def cleanup_sma_cell(value: Any) -> str:
    text = str(value or "").strip()
    if text.startswith('="') and text.endswith('"'):
        text = text[2:-1]
    return text.strip().strip('"').strip()


def parse_sma_portal_day(value: Any, year: int, month: int, field: str) -> dt.date:
    text = cleanup_sma_cell(value)
    match = re.match(r"^(\d{1,2})\.(\d{1,2})\.?(?:\d{2,4})?$", text)
    if not match:
        raise ValueError(f"invalid SMA portal day in {field}: {value!r}")
    day = int(match.group(1))
    parsed_month = int(match.group(2))
    if parsed_month != month:
        raise ValueError(f"SMA portal day {value!r} does not match filename month {month:02d}")
    return dt.date(year, month, day)


def sma_portal_component_name(header: str) -> str:
    return cleanup_sma_cell(header).split("/", 1)[0].strip()


def parse_date(value: Any, field: str = "date") -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        base = dt.datetime(1899, 12, 30)
        return (base + dt.timedelta(days=float(value))).date()
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"{field} is required")
    text = text.replace("/", "-")
    if "T" in text:
        text = text.split("T", 1)[0]
    if " " in text:
        text = text.split(" ", 1)[0]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%Y.%m.%d"):
        try:
            return dt.datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    raise ValueError(f"invalid {field}: {value!r}; expected YYYY-MM-DD or DD.MM.YYYY")


def parse_float(value: Any, field: str) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if math.isfinite(float(value)):
            return float(value)
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("\u00a0", " ").replace(" ", "")
    text = text.replace("kWh", "").replace("KWh", "").replace("Wh", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        value_float = float(text)
    except ValueError as exc:
        raise ValueError(f"invalid {field}: {value!r}") from exc
    if not math.isfinite(value_float):
        return None
    return value_float


def detect_delimiter(path: Path) -> str:
    sample = path.read_text(encoding="utf-8-sig", errors="replace")[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        return ";" if sample.count(";") > sample.count(",") else ","


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    delimiter = detect_delimiter(path)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=delimiter)
        if not reader.fieldnames:
            return []
        rows: list[dict[str, Any]] = []
        for raw in reader:
            rows.append({str(k or "").strip(): v for k, v in raw.items()})
    return rows


def read_xlsx_rows(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise SystemExit("Reading .xlsx SMA exports requires openpyxl. Use CSV or install openpyxl.") from exc
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet:
        if sheet not in wb.sheetnames:
            raise SystemExit(f"Sheet not found in {path}: {sheet}")
        ws = wb[sheet]
    else:
        ws = wb[wb.sheetnames[0]]
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw = {headers[idx]: value for idx, value in enumerate(row) if idx < len(headers)}
        if any(value not in (None, "") for value in raw.values()):
            rows.append(raw)
    return rows


def read_table(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_xlsx_rows(path, sheet)
    return read_csv_rows(path)


def find_column(headers: Iterable[str], requested: str | None, aliases: Iterable[str]) -> str | None:
    header_list = list(headers)
    normalized = {normalize_header(header): header for header in header_list}
    if requested:
        requested_norm = normalize_header(requested)
        if requested_norm in normalized:
            return normalized[requested_norm]
    for alias in aliases:
        alias_norm = normalize_header(alias)
        if alias_norm in normalized:
            return normalized[alias_norm]
    return None


def read_mapping(path: Path | None) -> dict[str, str]:
    if not path:
        return {}
    rows = read_table(path)
    mapping: dict[str, str] = {}
    for idx, row in enumerate(rows, start=2):
        sma_col = find_column(row.keys(), None, ["sma_name", "sma", "source", "portal_name", "name"])
        title_col = find_column(row.keys(), None, ["evcc_title", "title", "evcc", "vm_title"])
        if not sma_col or not title_col:
            raise SystemExit(f"Mapping file {path} must contain sma_name and evcc_title columns")
        sma_name = str(row.get(sma_col) or "").strip()
        title = str(row.get(title_col) or "").strip()
        if not sma_name and not title:
            continue
        if not sma_name or not title:
            raise SystemExit(f"Incomplete mapping row {idx} in {path}")
        mapping[sma_name] = title
        normalized_name = normalize_sma_name(sma_name)
        if normalized_name:
            mapping[normalized_name] = title
    return mapping


def infer_unit(column_name: str, explicit_unit: str) -> str:
    if explicit_unit != "auto":
        return explicit_unit
    header = normalize_header(column_name)
    if "wh" in header and "kwh" not in header:
        return "wh"
    return "kwh"


def to_wh(value: float, unit: str) -> float:
    return value if unit == "wh" else value * 1000.0


def detect_format(rows: list[dict[str, Any]], requested: str, date_column: str | None, name_column: str | None, energy_column: str | None) -> str:
    if requested != "auto":
        return requested
    if not rows:
        return "long"
    headers = rows[0].keys()
    found_name = find_column(headers, name_column, ["sma_name", "sma", "source", "portal_name", "name"])
    found_energy = find_column(headers, energy_column, ["energy_kwh", "energy_wh", "yield_kwh", "yield_wh", "pv_energy_kwh", "pv_energy_wh"])
    found_date = find_column(headers, date_column, ["date", "day", "datum", "time", "timestamp"])
    if found_date and found_name and found_energy:
        return "long"
    return "wide"


def map_title(sma_name: str, mapping: dict[str, str], require_mapping: bool) -> tuple[str, bool]:
    if sma_name in mapping:
        return mapping[sma_name], True
    normalized_name = normalize_sma_name(sma_name)
    if normalized_name in mapping:
        return mapping[normalized_name], True
    if require_mapping:
        raise SystemExit(f"Missing SMA mapping for {sma_name!r}")
    return sma_name, False


def should_exclude_sma_name(raw_name: str, normalized_name: str, exclude_patterns: list[re.Pattern[str]]) -> bool:
    return any(pattern.search(raw_name) or pattern.search(normalized_name) for pattern in exclude_patterns)


def parse_sma_portal_month_from_name(path: Path) -> tuple[int, int]:
    match = re.search(r"Analyse_(\d{4})_(\d{2})", path.name)
    if not match:
        raise SystemExit(f"SMA portal analysis file name must match Analyse_YYYY_MM.csv: {path}")
    return int(match.group(1)), int(match.group(2))


def parse_sma_portal_classic_analysis_file(path: Path, mapping: dict[str, str], require_mapping: bool, exclude_patterns: list[re.Pattern[str]]) -> tuple[list[SmaEnergyRow], set[str]]:
    year, month = parse_sma_portal_month_from_name(path)
    parsed: list[SmaEnergyRow] = []
    unmapped: set[str] = set()
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=";")
        try:
            header = next(reader)
        except StopIteration:
            return parsed, unmapped
        components: list[tuple[int, str, str]] = []
        for index, raw_header in enumerate(header[1:], start=1):
            raw_name = sma_portal_component_name(raw_header)
            if not raw_name:
                continue
            sma_name = normalize_sma_name(raw_name)
            if should_exclude_sma_name(raw_name, sma_name, exclude_patterns):
                continue
            components.append((index, raw_name, sma_name))
        if not components:
            raise SystemExit(f"No SMA portal component columns found in {path}")
        for row_number, row in enumerate(reader, start=2):
            if not row or not cleanup_sma_cell(row[0]):
                continue
            day = parse_sma_portal_day(row[0], year, month, f"{path.name} row {row_number}")
            for index, _raw_name, sma_name in components:
                raw_value = row[index] if index < len(row) else ""
                value_kwh = parse_float(cleanup_sma_cell(raw_value), f"{path.name} row {row_number} {sma_name}")
                if value_kwh is None:
                    continue
                title, mapped = map_title(sma_name, mapping, require_mapping)
                if not mapped:
                    unmapped.add(sma_name)
                parsed.append(SmaEnergyRow(day, sma_name, title, to_wh(value_kwh, "kwh")))
    return parsed, unmapped


def read_sma_portal_classic_analysis_dir(input_dir: Path, mapping: dict[str, str], require_mapping: bool, exclude_name_regex: list[str]) -> tuple[list[SmaEnergyRow], set[str]]:
    patterns = [re.compile(pattern) for pattern in exclude_name_regex if pattern]
    files = sorted(input_dir.rglob("Analyse_????_??.csv"))
    if not files:
        raise SystemExit(f"No SMA portal analysis files found below {input_dir}; expected Analyse_YYYY_MM.csv")
    rows: list[SmaEnergyRow] = []
    unmapped: set[str] = set()
    for path in files:
        file_rows, file_unmapped = parse_sma_portal_classic_analysis_file(path, mapping, require_mapping, patterns)
        rows.extend(file_rows)
        unmapped.update(file_unmapped)
    return rows, unmapped


def parse_long_rows(rows: list[dict[str, Any]], mapping: dict[str, str], require_mapping: bool, date_column: str | None, name_column: str | None, energy_column: str | None, energy_unit: str) -> tuple[list[SmaEnergyRow], set[str]]:
    parsed: list[SmaEnergyRow] = []
    unmapped: set[str] = set()
    for idx, row in enumerate(rows, start=2):
        date_col = find_column(row.keys(), date_column, ["date", "day", "datum", "time", "timestamp"])
        name_col = find_column(row.keys(), name_column, ["sma_name", "sma", "source", "portal_name", "name"])
        value_col = find_column(row.keys(), energy_column, ["energy_kwh", "energy_wh", "yield_kwh", "yield_wh", "pv_energy_kwh", "pv_energy_wh"])
        if not date_col or not name_col or not value_col:
            raise SystemExit("Long SMA input requires date, sma_name, and energy columns")
        sma_name = str(row.get(name_col) or "").strip()
        if not sma_name:
            continue
        raw_value = parse_float(row.get(value_col), f"{value_col} row {idx}")
        if raw_value is None:
            continue
        title, mapped = map_title(sma_name, mapping, require_mapping)
        if not mapped:
            unmapped.add(sma_name)
        unit = infer_unit(value_col, energy_unit)
        parsed.append(SmaEnergyRow(parse_date(row.get(date_col), date_col), sma_name, title, to_wh(raw_value, unit)))
    return parsed, unmapped


def parse_wide_rows(rows: list[dict[str, Any]], mapping: dict[str, str], require_mapping: bool, date_column: str | None, energy_unit: str) -> tuple[list[SmaEnergyRow], set[str]]:
    parsed: list[SmaEnergyRow] = []
    unmapped: set[str] = set()
    if not rows:
        return parsed, unmapped
    date_col = find_column(rows[0].keys(), date_column, ["date", "day", "datum", "time", "timestamp"])
    if not date_col:
        raise SystemExit("Wide SMA input requires a date column")
    value_columns = [header for header in rows[0].keys() if header != date_col and str(header).strip()]
    if not value_columns:
        raise SystemExit("Wide SMA input requires at least one SMA value column")
    for idx, row in enumerate(rows, start=2):
        day = parse_date(row.get(date_col), date_col)
        for column in value_columns:
            sma_name = str(column).strip()
            raw_value = parse_float(row.get(column), f"{column} row {idx}")
            if raw_value is None:
                continue
            title, mapped = map_title(sma_name, mapping, require_mapping)
            if not mapped:
                unmapped.add(sma_name)
            unit = infer_unit(column, energy_unit)
            parsed.append(SmaEnergyRow(day, sma_name, title, to_wh(raw_value, unit)))
    return parsed, unmapped


def read_sma_energy(path: Path, mapping: dict[str, str], require_mapping: bool, fmt: str, date_column: str | None, name_column: str | None, energy_column: str | None, energy_unit: str, sheet: str | None = None) -> tuple[list[SmaEnergyRow], set[str]]:
    rows = read_table(path, sheet)
    detected = detect_format(rows, fmt, date_column, name_column, energy_column)
    if detected == "long":
        return parse_long_rows(rows, mapping, require_mapping, date_column, name_column, energy_column, energy_unit)
    if detected == "wide":
        return parse_wide_rows(rows, mapping, require_mapping, date_column, energy_unit)
    raise SystemExit(f"Unsupported input format: {detected}")


def timestamp_for_day(day: dt.date, tz: ZoneInfo) -> int:
    local_noon = dt.datetime.combine(day, dt.time(hour=12), tzinfo=tz)
    return int(local_noon.astimezone(dt.timezone.utc).timestamp() * 1000)


def escape_label(value: str) -> str:
    return value.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n")


def line(metric: str, labels: dict[str, str], samples: list[tuple[int, float]]) -> str:
    metric_labels = {"__name__": metric, **labels}
    return json.dumps({
        "metric": metric_labels,
        "timestamps": [ts for ts, _ in samples],
        "values": [value for _, value in samples],
    }, separators=(",", ":"), ensure_ascii=False)


def build_series(rows: list[SmaEnergyRow], metric: str, tz: ZoneInfo, start_day: dt.date | None, end_day: dt.date | None) -> tuple[dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]], dict[str, Any]]:
    grouped_values: dict[tuple[dt.date, str], float] = {}
    sma_names_by_title: dict[str, set[str]] = {}
    for row in rows:
        if start_day and row.day < start_day:
            continue
        if end_day and row.day >= end_day:
            continue
        if row.energy_wh < 0:
            raise SystemExit(f"Negative SMA energy is not supported: {row.sma_name} {row.day} {row.energy_wh}")
        key = (row.day, row.title)
        grouped_values[key] = grouped_values.get(key, 0.0) + row.energy_wh
        sma_names_by_title.setdefault(row.title, set()).add(row.sma_name)

    series: dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]] = {}
    summary_by_title: dict[str, dict[str, Any]] = {}
    first_day: dt.date | None = None
    last_day: dt.date | None = None
    for (day, title), energy_wh in sorted(grouped_values.items()):
        if energy_wh <= 0:
            continue
        labels = {
            "title": title,
            "local_year": f"{day.year:04d}",
            "local_month": f"{day.month:02d}",
        }
        key = (metric, tuple(sorted(labels.items())))
        series.setdefault(key, []).append((timestamp_for_day(day, tz), float(energy_wh)))
        item = summary_by_title.setdefault(title, {"title": title, "sma_names": set(), "days": 0, "energy_kwh": 0.0})
        item["sma_names"].update(sma_names_by_title.get(title, set()))
        item["days"] += 1
        item["energy_kwh"] += energy_wh / 1000.0
        first_day = day if first_day is None else min(first_day, day)
        last_day = day if last_day is None else max(last_day, day)

    title_summary = []
    for item in sorted(summary_by_title.values(), key=lambda value: value["title"]):
        title_summary.append({
            "title": item["title"],
            "sma_names": sorted(item["sma_names"]),
            "days": item["days"],
            "energy_kwh": round(item["energy_kwh"], 6),
        })
    summary = {
        "metric": metric,
        "series": len(series),
        "samples": sum(len(samples) for samples in series.values()),
        "first_day": first_day.isoformat() if first_day else None,
        "last_day": last_day.isoformat() if last_day else None,
        "titles": title_summary,
    }
    return series, summary


def post_bytes(url: str, data: bytes, timeout: int = 120) -> None:
    request = urllib.request.Request(url, data=data, method="POST", headers={"Content-Type": "application/jsonl"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} for {url}: {body}") from exc


def delete_metric(base_url: str, metric: str) -> None:
    params = urllib.parse.urlencode({"match[]": metric})
    url = f"{base_url.rstrip('/')}/api/v1/admin/tsdb/delete_series?{params}"
    request = urllib.request.Request(url, data=b"", method="POST")
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} while deleting {metric}: {body}") from exc


def import_series(base_url: str, series: dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]], batch_size: int) -> int:
    batch: list[str] = []
    batches = 0
    for (metric, label_items), samples in series.items():
        batch.append(line(metric, dict(label_items), sorted(samples)))
        if len(batch) >= batch_size:
            post_bytes(f"{base_url.rstrip('/')}/api/v1/import", ("\n".join(batch) + "\n").encode("utf-8"))
            batches += 1
            batch = []
    if batch:
        post_bytes(f"{base_url.rstrip('/')}/api/v1/import", ("\n".join(batch) + "\n").encode("utf-8"))
        batches += 1
    return batches


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import SMA PV daily energy exports into VictoriaMetrics")
    parser.add_argument("--vm-base-url", default="http://127.0.0.1:8428")
    parser.add_argument("--input-file", help="SMA CSV/XLSX export file")
    parser.add_argument("--input-dir", help="Directory containing SMA Portal Classic Analyse_YYYY_MM.csv files")
    parser.add_argument("--map-file", help="CSV/XLSX mapping file with sma_name and evcc_title columns")
    parser.add_argument("--format", choices=["auto", "long", "wide", "sma-portal-classic-analysis"], default="auto")
    parser.add_argument("--date-column", help="date column name; auto-detected when omitted")
    parser.add_argument("--name-column", help="SMA name column for long format; auto-detected when omitted")
    parser.add_argument("--energy-column", help="energy column for long format; auto-detected when omitted")
    parser.add_argument("--energy-unit", choices=["auto", "wh", "kwh"], default="auto")
    parser.add_argument("--sheet", help="Excel sheet name; defaults to first sheet")
    parser.add_argument("--metric", default=DEFAULT_METRIC, help="target metric for per-title daily Wh values; defaults to EVCC-compatible evcc_pv_energy_by_title_daily_wh")
    parser.add_argument("--timezone", default="Europe/Berlin")
    parser.add_argument("--start", help="first local day, YYYY-MM-DD")
    parser.add_argument("--end", help="exclusive local end day, YYYY-MM-DD")
    parser.add_argument("--require-mapping", action="store_true", help="fail when an SMA source name has no mapping to an EVCC title")
    parser.add_argument("--exclude-name-regex", action="append", default=[], help="regex for SMA source names to skip, e.g. portal aggregate totals")
    parser.add_argument("--replace", action="store_true", help="delete the target metric before writing")
    parser.add_argument("--write", action="store_true", help="write imported metrics to VictoriaMetrics")
    parser.add_argument("--batch-size", type=int, default=200)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    started = time.perf_counter()
    log(f"import-sma-pv-energy.py v{SCRIPT_VERSION} (last modified {SCRIPT_LAST_MODIFIED})")
    map_path = Path(args.map_file) if args.map_file else None
    if map_path and not map_path.exists():
        raise SystemExit(f"SMA mapping file not found: {map_path}")
    mapping = read_mapping(map_path)
    if args.format == "sma-portal-classic-analysis":
        if not args.input_dir:
            raise SystemExit("--input-dir is required with --format=sma-portal-classic-analysis")
        input_dir = Path(args.input_dir)
        if not input_dir.exists():
            raise SystemExit(f"SMA input directory not found: {input_dir}")
        rows, unmapped = read_sma_portal_classic_analysis_dir(input_dir, mapping, args.require_mapping, args.exclude_name_regex)
        input_label = str(input_dir)
    else:
        if not args.input_file:
            raise SystemExit("--input-file is required unless --format=sma-portal-classic-analysis is used")
        input_path = Path(args.input_file)
        if not input_path.exists():
            raise SystemExit(f"SMA input file not found: {input_path}")
        rows, unmapped = read_sma_energy(
            input_path,
            mapping,
            args.require_mapping,
            args.format,
            args.date_column,
            args.name_column,
            args.energy_column,
            args.energy_unit,
            args.sheet,
        )
        input_label = str(input_path)
    if not rows:
        raise SystemExit(f"No SMA energy rows found in {input_label}")
    tz = ZoneInfo(args.timezone)
    start_day = parse_date(args.start, "--start") if args.start else None
    end_day = parse_date(args.end, "--end") if args.end else None
    if start_day and end_day and end_day <= start_day:
        raise SystemExit("--end must be after --start")
    series, summary = build_series(rows, args.metric, tz, start_day, end_day)
    log(json.dumps(summary, indent=2, ensure_ascii=False))
    if unmapped:
        log("Unmapped SMA names used as title: " + ", ".join(sorted(unmapped)))
        log("Use --require-mapping to fail instead, or provide --map-file with sma_name,evcc_title.")
    if args.write:
        if args.replace:
            log(f"Delete existing metric: {args.metric}")
            delete_metric(args.vm_base_url, args.metric)
        batches = import_series(args.vm_base_url, series, args.batch_size)
        log(f"Imported batches: {batches}")
    else:
        log("Dry run only. Re-run with --write to import metrics.")
    log(f"Finished in {time.perf_counter() - started:.1f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

