#!/usr/bin/env python3
"""
Script: import-investment-costs.py
Purpose: Calculate PV investment cost rollups from a local investment file and VictoriaMetrics PV data.
Version: 2026.06.11.1
Last modified: 2026-06-11
"""
from __future__ import annotations

import argparse
import ast
import csv
import datetime as dt
import json
import math
import operator
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import warnings
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

SCRIPT_VERSION = "2026.06.11.1"
SCRIPT_LAST_MODIFIED = "2026-06-11"
GENERATED_METRICS = [
    "evcc_pv_investment_cost_daily_eur",
    "evcc_pv_lcoe_daily_ct_per_kwh",
    "evcc_pv_lcoe_energy_monthly_wh",
    "evcc_pv_investment_cost_monthly_eur",
    "evcc_pv_lcoe_cost_monthly_eur",
    "evcc_pv_lcoe_monthly_ct_per_kwh",
    "evcc_pv_lcoe_yearly_ct_per_kwh",
    "evcc_pv_lcoe_period_ct_per_kwh",
    "evcc_pv_lcoe_rolling_7d_ct_per_kwh",
    "evcc_pv_installed_watt_peak_yearly",
    "evcc_pv_energy_by_title_yearly_wh",
    "evcc_pv_energy_by_title_yearly_with_coverage_wh",
    "evcc_pv_specific_yield_yearly_kwh_per_kwp",
    "evcc_pv_specific_yield_yearly_with_coverage_kwh_per_kwp",
    "evcc_pv_specific_yield_rolling_7d_kwh_per_kwp",
    "evcc_pv_effective_lcoe_daily_ct_per_kwh",
    "evcc_pv_effective_lcoe_yearly_ct_per_kwh",
    "evcc_pv_effective_lcoe_period_ct_per_kwh",
    "evcc_pv_effective_lcoe_monthly_ct_per_kwh",
]

LEGACY_GENERATED_METRICS = [
    "evcc_pv_lcoe_coverage_ratio",
    "evcc_pv_lcoe_partial",
]

OPS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.Pow: operator.pow,
    ast.USub: operator.neg,
    ast.UAdd: operator.pos,
}


def log(message: str) -> None:
    print(message, flush=True)


def safe_formula_value(value: Any) -> Any:
    if not isinstance(value, str) or not value.startswith("="):
        return value
    expr = value[1:].strip()
    if not expr or "[" in expr or "]" in expr or "!" in expr or ":" in expr:
        return None
    try:
        node = ast.parse(expr, mode="eval").body
    except SyntaxError:
        return None

    def eval_node(n: ast.AST) -> float:
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return float(n.value)
        if isinstance(n, ast.BinOp) and type(n.op) in OPS:
            return float(OPS[type(n.op)](eval_node(n.left), eval_node(n.right)))
        if isinstance(n, ast.UnaryOp) and type(n.op) in OPS:
            return float(OPS[type(n.op)](eval_node(n.operand)))
        raise ValueError(ast.dump(n))

    try:
        return eval_node(node)
    except Exception:
        return None


def parse_date(value: Any, field: str, row_name: str) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        base = dt.datetime(1899, 12, 30)
        return (base + dt.timedelta(days=float(value))).date()
    if isinstance(value, str):
        value = value.strip()
        if not value:
            raise ValueError(f"{row_name}: {field} is required")
        try:
            return dt.date.fromisoformat(value[:10])
        except ValueError as exc:
            raise ValueError(f"{row_name}: invalid {field}: {value!r}; expected YYYY-MM-DD") from exc
    raise ValueError(f"{row_name}: invalid {field}: {value!r}")


def parse_float(value: Any, field: str, row_name: str, required: bool = True) -> float | None:
    value = safe_formula_value(value)
    if value is None or value == "":
        if required:
            raise ValueError(f"{row_name}: {field} is required")
        return None
    if isinstance(value, str):
        value = value.strip().replace("€", "").replace(" ", "").replace(",", ".")
        if "/" in value:
            parts = value.split("/", 1)
            try:
                numerator = float(parts[0])
                denominator = float(parts[1])
                if denominator != 0:
                    return numerator / denominator
            except Exception:
                pass
    try:
        return float(value)
    except Exception as exc:
        raise ValueError(f"{row_name}: invalid {field}: {value!r}") from exc

def is_yes(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "y", "on"}


def normalize_row(raw: dict[str, Any], row_number: int) -> dict[str, Any] | None:
    if not any(v not in (None, "") for v in raw.values()):
        return None
    row_name = str(raw.get("asset_id") or f"row {row_number}")
    asset_id = str(raw.get("asset_id") or "").strip()
    asset_type = str(raw.get("asset_type") or "").strip().lower()
    evcc_title = str(raw.get("evcc_title") or "").strip()
    include_raw = raw.get("include_in_effective_price")
    include_text = str(include_raw or "").strip().lower()
    include = is_yes(include_text)

    key_values = {
        "asset_id": asset_id,
        "asset_type": asset_type,
        "evcc_title": evcc_title,
    }
    if not any(key_values.values()):
        log(f"Skip investment row {row_number}: no asset_id, asset_type, or evcc_title")
        return None
    missing = [key for key, value in key_values.items() if not value]
    if missing:
        raise ValueError(f"{row_name}: missing mandatory field(s): {', '.join(missing)}")

    base_asset = {
        "asset_id": asset_id,
        "asset_type": asset_type,
        "evcc_title": evcc_title,
        "include": include,
        "notes": str(raw.get("notes") or "").strip(),
    }

    if not include_text:
        log(f"Skip investment row: {asset_id} / {evcc_title} has no include_in_effective_price value")
        return {**base_asset, "include": False}
    if asset_type not in {"pv", "pv_shared"}:
        log(f"Skip investment row for PV cost import: {asset_id} has asset_type={asset_type}")
        return base_asset
    if not include:
        return base_asset

    required_fields = ["commissioning_date", "purchase_price_eur", "lifetime_years"]
    if asset_type == "pv":
        required_fields.append("watt_peak")
    if asset_type == "pv_shared":
        required_fields.append("allocation_percent")
    missing_fields = []
    for field in required_fields:
        raw_value = raw.get(field)
        if raw_value in (None, "") or safe_formula_value(raw_value) is None:
            missing_fields.append(field)
    if missing_fields:
        log(f"Skip incomplete included PV cost row: {asset_id} / {evcc_title} misses {', '.join(missing_fields)}")
        return {**base_asset, "include": False}

    commissioning_date = parse_date(raw.get("commissioning_date"), "commissioning_date", row_name)
    purchase_price_eur = parse_float(raw.get("purchase_price_eur"), "purchase_price_eur", row_name)
    lifetime_years = parse_float(raw.get("lifetime_years"), "lifetime_years", row_name)
    yearly_opex_eur = parse_float(raw.get("yearly_opex_eur"), "yearly_opex_eur", row_name, required=False) or 0.0
    watt_peak = parse_float(raw.get("watt_peak"), "watt_peak", row_name, required=False)
    capacity_wh = parse_float(raw.get("capacity_wh"), "capacity_wh", row_name, required=False)
    allocation_percent = parse_float(raw.get("allocation_percent"), "allocation_percent", row_name, required=False)
    if allocation_percent is None:
        allocation_percent = 1.0

    if lifetime_years is None or lifetime_years <= 0:
        raise ValueError(f"{row_name}: lifetime_years must be > 0")
    if purchase_price_eur is None or purchase_price_eur < 0:
        raise ValueError(f"{row_name}: purchase_price_eur must be >= 0")
    if allocation_percent < 0:
        raise ValueError(f"{row_name}: allocation_percent must be >= 0")
    if asset_type == "pv" and not watt_peak:
        raise ValueError(f"{row_name}: watt_peak is required for asset_type=pv")

    return {
        **base_asset,
        "commissioning_date": commissioning_date,
        "purchase_price_eur": purchase_price_eur,
        "lifetime_years": lifetime_years,
        "yearly_opex_eur": yearly_opex_eur,
        "watt_peak": watt_peak,
        "capacity_wh": capacity_wh,
        "allocation_percent": allocation_percent,
    }

def read_csv(path: Path) -> list[dict[str, Any]]:
    rows = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for idx, row in enumerate(reader, start=2):
            normalized = normalize_row(row, idx)
            if normalized:
                rows.append(normalized)
    return rows


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise SystemExit("Reading .xlsx investment files requires the optional Python package openpyxl. Use CSV or install openpyxl.") from exc
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
        formula_wb = openpyxl.load_workbook(path, data_only=False)
        value_wb = openpyxl.load_workbook(path, data_only=True)
    if "Investments" not in formula_wb.sheetnames or "Investments" not in value_wb.sheetnames:
        raise SystemExit(f"{path} must contain a sheet named Investments")
    formula_ws = formula_wb["Investments"]
    value_ws = value_wb["Investments"]
    headers = [str(cell.value or "").strip() for cell in formula_ws[1]]
    rows = []
    for row_number in range(2, formula_ws.max_row + 1):
        raw = {}
        for index, header in enumerate(headers, start=1):
            value = value_ws.cell(row=row_number, column=index).value
            if value in (None, ""):
                value = formula_ws.cell(row=row_number, column=index).value
            raw[header] = value
        normalized = normalize_row(raw, row_number)
        if normalized:
            rows.append(normalized)
    return rows


def read_investments(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        return read_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise SystemExit(f"Unsupported investment file type: {path.suffix}. Use .xlsx or .csv")


def http_json(url: str, timeout: int = 120) -> dict[str, Any]:
    try:
        with urllib.request.urlopen(url, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} for {url}: {body}") from exc


def post_bytes(url: str, data: bytes, timeout: int = 120) -> None:
    request = urllib.request.Request(url, data=data, method="POST", headers={"Content-Type": "application/jsonl"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} for {url}: {body}") from exc


def escape_label(value: str) -> str:
    return value.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n")


def query_range(base_url: str, query: str, start: dt.datetime, end: dt.datetime, step: str) -> dict[str, Any]:
    params = urllib.parse.urlencode({
        "query": query,
        "start": start.isoformat().replace("+00:00", "Z"),
        "end": end.isoformat().replace("+00:00", "Z"),
        "step": step,
    })
    result = http_json(f"{base_url.rstrip('/')}/api/v1/query_range?{params}")
    if result.get("status") != "success":
        raise RuntimeError(f"query_range failed: {result}")
    return result


def delete_series(base_url: str, selector: str) -> None:
    params = urllib.parse.urlencode({"match[]": selector})
    url = f"{base_url.rstrip('/')}/api/v1/admin/tsdb/delete_series?{params}"
    request = urllib.request.Request(url, data=b"", method="POST")
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} while deleting {selector}: {body}") from exc


def delete_metric(base_url: str, metric: str) -> None:
    delete_series(base_url, metric)


def pv_energy_rollup_source_label(energy_source: str) -> str:
    return "sma" if energy_source == "daily-metric" else ("combined" if energy_source == "combined" else "helper")

def active_until(asset: dict[str, Any]) -> dt.date:
    start = asset["commissioning_date"]
    years = int(math.ceil(float(asset["lifetime_years"])))
    try:
        return start.replace(year=start.year + years)
    except ValueError:
        return start.replace(month=2, day=28, year=start.year + years)


def daily_cost(asset: dict[str, Any], day: dt.date) -> float:
    if day < asset["commissioning_date"] or day >= active_until(asset):
        return 0.0
    base_cost = (asset["purchase_price_eur"] / (asset["lifetime_years"] * 365.25)) + (asset["yearly_opex_eur"] / 365.25)
    return base_cost * float(asset.get("allocation_percent", 1.0))

def local_date_from_eval(ts: float, tz: ZoneInfo) -> dt.date:
    evaluated = dt.datetime.fromtimestamp(float(ts), tz=dt.timezone.utc).astimezone(tz).date()
    return evaluated - dt.timedelta(days=1)


def local_date_from_same_day_eval(ts: float, tz: ZoneInfo) -> dt.date:
    return dt.datetime.fromtimestamp(float(ts), tz=dt.timezone.utc).astimezone(tz).date()


def validate_metric_name(metric: str) -> str:
    if not re.match(r"^[a-zA-Z_:][a-zA-Z0-9_:]*$", metric):
        raise ValueError(f"invalid metric name: {metric!r}")
    return metric

def date_chunks(start_day: dt.date, end_day: dt.date, chunk_days: int) -> Iterable[tuple[dt.date, dt.date]]:
    current = start_day
    while current < end_day:
        next_day = min(current + dt.timedelta(days=chunk_days), end_day)
        yield current, next_day
        current = next_day


def fetch_daily_energy(base_url: str, asset: dict[str, Any], start_day: dt.date, end_day: dt.date, tz: ZoneInfo, peak_limit: float, sample_interval: str) -> dict[dt.date, float]:
    title = escape_label(asset["evcc_title"])
    expr = f'sum(integrate(((pvPower_value{{title="{title}"}} < {peak_limit:g}) default 0)[1d:{sample_interval}]) / 3600)'
    values: dict[dt.date, float] = {}
    for chunk_start, chunk_end in date_chunks(start_day, end_day, 30):
        eval_start_local = dt.datetime.combine(chunk_start + dt.timedelta(days=1), dt.time.min, tzinfo=tz)
        eval_end_local = dt.datetime.combine(chunk_end, dt.time.min, tzinfo=tz)
        if eval_end_local <= eval_start_local:
            continue
        data = query_range(
            base_url,
            expr,
            eval_start_local.astimezone(dt.timezone.utc),
            eval_end_local.astimezone(dt.timezone.utc),
            "1d",
        )
        for series in data.get("data", {}).get("result", []):
            for ts, value in series.get("values", []):
                day = local_date_from_eval(float(ts), tz)
                if day < start_day or day >= end_day:
                    continue
                try:
                    values[day] = values.get(day, 0.0) + float(value)
                except (TypeError, ValueError):
                    continue
    return values


def fetch_daily_energy_from_daily_metric(base_url: str, asset: dict[str, Any], start_day: dt.date, end_day: dt.date, tz: ZoneInfo, metric: str) -> dict[dt.date, float]:
    validate_metric_name(metric)
    title = escape_label(asset["evcc_title"])
    selector = f'{metric}{{title="{title}"}}'
    start_local = dt.datetime.combine(start_day, dt.time.min, tzinfo=tz)
    end_local = dt.datetime.combine(end_day, dt.time.min, tzinfo=tz)
    params = urllib.parse.urlencode({
        "match[]": selector,
        "start": start_local.astimezone(dt.timezone.utc).isoformat().replace("+00:00", "Z"),
        "end": end_local.astimezone(dt.timezone.utc).isoformat().replace("+00:00", "Z"),
    })
    url = f"{base_url.rstrip('/')}/api/v1/export?{params}"
    values: dict[dt.date, float] = {}
    try:
        with urllib.request.urlopen(url, timeout=120) as response:
            for raw_line in response:
                if not raw_line.strip():
                    continue
                item = json.loads(raw_line.decode("utf-8"))
                for timestamp_ms, value in zip(item.get("timestamps", []), item.get("values", [])):
                    day = dt.datetime.fromtimestamp(float(timestamp_ms) / 1000.0, tz=dt.timezone.utc).astimezone(tz).date()
                    if day < start_day or day >= end_day:
                        continue
                    try:
                        energy_wh = float(value)
                    except (TypeError, ValueError):
                        continue
                    if energy_wh > 0:
                        values[day] = values.get(day, 0.0) + energy_wh
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} while exporting {selector}: {body}") from exc
    return values


def merge_daily_energy(evcc_values: dict[dt.date, float], metric_values: dict[dt.date, float], conflict_policy: str) -> tuple[dict[dt.date, float], dict[str, int]]:
    merged = dict(evcc_values)
    stats = {
        "evcc_days": len([value for value in evcc_values.values() if value > 0]),
        "metric_days": len([value for value in metric_values.values() if value > 0]),
        "overlap_days": 0,
        "merged_days": 0,
    }
    for day, metric_wh in sorted(metric_values.items()):
        if metric_wh <= 0:
            continue
        evcc_wh = merged.get(day, 0.0)
        if evcc_wh > 0:
            stats["overlap_days"] += 1
            if conflict_policy == "prefer-evcc":
                continue
            if conflict_policy == "prefer-daily-metric":
                merged[day] = metric_wh
                continue
            if conflict_policy == "sum":
                merged[day] = evcc_wh + metric_wh
                continue
            if conflict_policy == "error":
                raise RuntimeError(f"Overlapping PV energy for {day}; choose another --combined-energy-conflict policy")
            raise ValueError(f"unsupported conflict policy: {conflict_policy}")
        merged[day] = metric_wh
    stats["merged_days"] = len([value for value in merged.values() if value > 0])
    return merged, stats

def line(metric: str, labels: dict[str, str], samples: list[tuple[int, float]]) -> str:
    metric_labels = {"__name__": metric, **labels}
    return json.dumps({
        "metric": metric_labels,
        "timestamps": [ts for ts, _ in samples],
        "values": [value for _, value in samples],
    }, separators=(",", ":"))


def append_sample(series: dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]], metric: str, labels: dict[str, str], timestamp_ms: int, value: float) -> None:
    if not math.isfinite(value):
        return
    key = (metric, tuple(sorted(labels.items())))
    series.setdefault(key, []).append((timestamp_ms, float(value)))


def timestamp_for_day(day: dt.date, tz: ZoneInfo) -> int:
    local_noon = dt.datetime.combine(day, dt.time(hour=12), tzinfo=tz)
    return int(local_noon.astimezone(dt.timezone.utc).timestamp() * 1000)


def timestamp_for_year(year: int, tz: ZoneInfo) -> int:
    local_noon = dt.datetime(year, 1, 1, 12, 0, tzinfo=tz)
    return int(local_noon.astimezone(dt.timezone.utc).timestamp() * 1000)

def timestamp_for_month(year: int, month: int, tz: ZoneInfo) -> int:
    local_noon = dt.datetime(year, month, 15, 12, 0, tzinfo=tz)
    return int(local_noon.astimezone(dt.timezone.utc).timestamp() * 1000)


def iter_days(start_day: dt.date, end_day: dt.date) -> Iterable[dt.date]:
    current = start_day
    while current < end_day:
        yield current
        current += dt.timedelta(days=1)


def active_cost_days_for_calendar_year(assets: list[dict[str, Any]], year: int) -> float:
    year_start = dt.date(year, 1, 1)
    year_end = dt.date(year + 1, 1, 1)
    expected_days = 0.0
    for day in iter_days(year_start, year_end):
        if any(daily_cost(asset, day) > 0 for asset in assets):
            expected_days += 1.0
    return expected_days

def installed_watt_peak_for_day(assets: list[dict[str, Any]], day: dt.date) -> float:
    total = 0.0
    for asset in assets:
        watt_peak = float(asset.get("watt_peak") or 0.0)
        if watt_peak <= 0 or day < asset["commissioning_date"]:
            continue
        total += watt_peak * float(asset.get("allocation_percent", 1.0))
    return total


def average_installed_watt_peak_for_range(assets: list[dict[str, Any]], start_day: dt.date, end_day: dt.date) -> float:
    total_watt_days = 0.0
    days = 0
    for day in iter_days(start_day, end_day):
        total_watt_days += installed_watt_peak_for_day(assets, day)
        days += 1
    return total_watt_days / days if days else 0.0


def average_installed_watt_peak_for_calendar_year(assets: list[dict[str, Any]], year: int) -> float:
    return average_installed_watt_peak_for_range(assets, dt.date(year, 1, 1), dt.date(year + 1, 1, 1))

def validate_shared_allocations(assets: list[dict[str, Any]], tolerance: float = 0.001) -> list[dict[str, Any]]:
    shared_groups: dict[str, list[dict[str, Any]]] = {}
    for asset in assets:
        if asset.get("asset_type") == "pv_shared" and asset.get("include"):
            shared_groups.setdefault(asset["asset_id"], []).append(asset)

    summaries = []
    for asset_id, rows in sorted(shared_groups.items()):
        allocation_sum = sum(float(row.get("allocation_percent", 0.0)) for row in rows)
        titles = sorted({row["evcc_title"] for row in rows})
        status = "ok" if abs(allocation_sum - 1.0) <= tolerance else "warning"
        if status == "warning":
            log(f"WARNING shared asset allocation does not sum to 100%: {asset_id} = {allocation_sum * 100:.3f}% across {', '.join(titles)}")
        else:
            log(f"Shared asset allocation OK: {asset_id} = {allocation_sum * 100:.3f}% across {', '.join(titles)}")
        summaries.append({
            "asset_id": asset_id,
            "titles": titles,
            "allocation_sum": allocation_sum,
            "status": status,
        })
    return summaries

def coverage_ratio(covered_days: float, expected_days: float) -> float | None:
    if expected_days <= 0:
        return None
    return max(0.0, min(1.0, covered_days / expected_days))



def should_write_lcoe(ratio: float | None, min_lcoe_coverage_ratio: float) -> bool:
    return ratio is None or ratio >= min_lcoe_coverage_ratio


def coverage_label(ratio: float | None) -> str:
    if ratio is None:
        return "n/a"
    return f"{round(ratio * 100):.0f}%"


def new_cost_bucket() -> dict[str, float]:
    return {
        "energy_wh": 0.0,
        "active_cost_eur": 0.0,
        "covered_cost_eur": 0.0,
        "expected_days": 0.0,
        "covered_days": 0.0,
    }


def add_bucket(target: dict[str, float], source: dict[str, float]) -> None:
    for key in ("energy_wh", "active_cost_eur", "covered_cost_eur", "expected_days", "covered_days"):
        target[key] = target.get(key, 0.0) + source.get(key, 0.0)


def build_rollups(base_url: str, assets: list[dict[str, Any]], start_day: dt.date, end_day: dt.date, tz: ZoneInfo, peak_limit: float, sample_interval: str, energy_source: str = "pv-power", pv_energy_metric: str = "evcc_pv_energy_by_title_daily_wh", skip_titles_without_energy: bool = False, write_pv_energy_rollup: bool = False, combined_energy_conflict: str = "prefer-evcc", min_lcoe_coverage_ratio: float = 0.0, partial_warning_threshold: float = 0.95) -> tuple[dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]], dict[str, Any]]:
    pv_assets = [a for a in assets if a["asset_type"] == "pv" and a["include"]]
    pv_cost_assets = [a for a in assets if a["asset_type"] in {"pv", "pv_shared"} and a["include"]]
    shared_allocation_summaries = validate_shared_allocations(assets)
    assets_by_title: dict[str, list[dict[str, Any]]] = {}
    for asset in pv_cost_assets:
        assets_by_title.setdefault(asset["evcc_title"], []).append(asset)

    series: dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]] = {}
    totals: dict[dt.date, dict[str, float]] = {}
    monthly_totals: dict[tuple[int, int], dict[str, float]] = {}
    asset_summaries = []
    title_summaries = []
    coverage_summaries = []

    for title, title_assets in sorted(assets_by_title.items()):
        if len(title_assets) == 1:
            log(f"Query PV energy: {title_assets[0]['asset_id']} / {title}")
        else:
            asset_ids = ", ".join(asset["asset_id"] for asset in title_assets)
            log(f"Query PV energy once for {title}: {len(title_assets)} investment rows ({asset_ids})")

        energy_merge_stats = None
        if energy_source == "daily-metric":
            energy_by_day = fetch_daily_energy_from_daily_metric(base_url, title_assets[0], start_day, end_day, tz, pv_energy_metric)
        elif energy_source == "pv-power":
            energy_by_day = fetch_daily_energy(base_url, title_assets[0], start_day, end_day, tz, peak_limit, sample_interval)
        elif energy_source == "combined":
            evcc_energy_by_day = fetch_daily_energy(base_url, title_assets[0], start_day, end_day, tz, peak_limit, sample_interval)
            metric_energy_by_day = fetch_daily_energy_from_daily_metric(base_url, title_assets[0], start_day, end_day, tz, pv_energy_metric)
            energy_by_day, energy_merge_stats = merge_daily_energy(evcc_energy_by_day, metric_energy_by_day, combined_energy_conflict)
            log(
                f"Combined energy for {title}: evcc_days={energy_merge_stats['evcc_days']}, "
                f"metric_days={energy_merge_stats['metric_days']}, overlap_days={energy_merge_stats['overlap_days']}, "
                f"merged_days={energy_merge_stats['merged_days']}, conflict_policy={combined_energy_conflict}"
            )
        else:
            raise ValueError(f"unsupported energy_source: {energy_source}")
        has_positive_energy = any(value > 0 for value in energy_by_day.values())
        if skip_titles_without_energy and not has_positive_energy:
            log(f"Skip title without positive energy samples: {title}")
            continue

        title_energy_wh = 0.0
        title_active_cost_eur = 0.0
        title_covered_cost_eur = 0.0
        title_expected_days = 0.0
        title_covered_days = 0.0
        monthly_title_totals: dict[tuple[int, int], dict[str, float]] = {}
        monthly_asset_costs: dict[tuple[str, tuple[int, int]], float] = {}
        asset_active_costs: dict[str, float] = {asset["asset_id"]: 0.0 for asset in title_assets}
        asset_covered_costs: dict[str, float] = {asset["asset_id"]: 0.0 for asset in title_assets}
        daily_title_values: dict[dt.date, dict[str, float]] = {}

        for day in iter_days(start_day, end_day):
            energy_wh = max(0.0, energy_by_day.get(day, 0.0))
            has_energy = energy_wh > 0
            month_key = (day.year, day.month)
            ts = timestamp_for_day(day, tz)
            energy_labels = {
                "title": title,
                "local_year": f"{day.year:04d}",
                "local_month": f"{day.month:02d}",
            }

            monthly_title_totals.setdefault(month_key, new_cost_bucket())
            daily_values = daily_title_values.setdefault(day, new_cost_bucket())

            if energy_wh > 0:
                monthly_title_totals[month_key]["energy_wh"] += energy_wh
                monthly_totals.setdefault(month_key, {"energy_wh": 0.0, "cost_eur": 0.0})
                monthly_totals[month_key]["energy_wh"] += energy_wh
                totals.setdefault(day, {"energy_wh": 0.0, "cost_eur": 0.0})
                totals[day]["energy_wh"] += energy_wh
                title_energy_wh += energy_wh
                daily_values["energy_wh"] += energy_wh

            daily_active_cost_eur = 0.0
            daily_asset_costs: list[tuple[dict[str, Any], float]] = []
            for asset in title_assets:
                cost_eur = daily_cost(asset, day)
                if cost_eur <= 0:
                    continue
                cost_labels = {
                    "asset_id": asset["asset_id"],
                    "title": title,
                    "local_year": f"{day.year:04d}",
                    "local_month": f"{day.month:02d}",
                }
                append_sample(series, "evcc_pv_investment_cost_daily_eur", cost_labels, ts, cost_eur)
                monthly_asset_costs[(asset["asset_id"], month_key)] = monthly_asset_costs.get((asset["asset_id"], month_key), 0.0) + cost_eur
                asset_active_costs[asset["asset_id"]] += cost_eur
                daily_asset_costs.append((asset, cost_eur))
                daily_active_cost_eur += cost_eur

            if daily_active_cost_eur > 0:
                monthly_title_totals[month_key]["active_cost_eur"] += daily_active_cost_eur
                monthly_title_totals[month_key]["expected_days"] += 1.0
                title_active_cost_eur += daily_active_cost_eur
                title_expected_days += 1.0
                daily_values["active_cost_eur"] += daily_active_cost_eur
                daily_values["expected_days"] = 1.0

            if daily_active_cost_eur > 0 and has_energy:
                monthly_title_totals[month_key]["covered_cost_eur"] += daily_active_cost_eur
                monthly_title_totals[month_key]["covered_days"] += 1.0
                monthly_totals.setdefault(month_key, {"energy_wh": 0.0, "cost_eur": 0.0})
                monthly_totals[month_key]["cost_eur"] += daily_active_cost_eur
                totals.setdefault(day, {"energy_wh": 0.0, "cost_eur": 0.0})
                totals[day]["cost_eur"] += daily_active_cost_eur
                title_covered_cost_eur += daily_active_cost_eur
                title_covered_days += 1.0
                daily_values["covered_cost_eur"] += daily_active_cost_eur
                daily_values["covered_days"] = 1.0
                for asset, cost_eur in daily_asset_costs:
                    asset_covered_costs[asset["asset_id"]] += cost_eur
                append_sample(series, "evcc_pv_lcoe_daily_ct_per_kwh", energy_labels, ts, daily_active_cost_eur / (energy_wh / 1000.0) * 100.0)

        yearly_title_totals: dict[int, dict[str, float]] = {}
        for (year, month), values in sorted(monthly_title_totals.items()):
            labels = {
                "title": title,
                "period": "month",
                "local_year": f"{year:04d}",
                "local_month": f"{month:02d}",
            }
            month_ratio = coverage_ratio(values["covered_days"], values["expected_days"])
            if month_ratio is not None and month_ratio < partial_warning_threshold:
                log(
                    f"WARNING partial PV energy coverage for {title} {year:04d}-{month:02d}: "
                    f"covered_days={int(values['covered_days'])}, expected_days={int(values['expected_days'])}, ratio={month_ratio:.3f}"
                )
            yearly_title_totals.setdefault(year, new_cost_bucket())
            add_bucket(yearly_title_totals[year], values)

        title_calendar_expected_days = 0.0
        for year, values in sorted(yearly_title_totals.items()):
            year_expected_days = active_cost_days_for_calendar_year(title_assets, year)
            title_calendar_expected_days += year_expected_days
            ratio = coverage_ratio(values["covered_days"], year_expected_days)
            if ratio is not None:
                coverage_summaries.append({
                    "title": title,
                    "period": "year",
                    "local_year": f"{year:04d}",
                    "covered_days": int(values["covered_days"]),
                    "expected_days": int(year_expected_days),
                    "coverage_ratio": ratio,
                    "partial": ratio < partial_warning_threshold,
                })
                if ratio < partial_warning_threshold:
                    log(
                        f"WARNING partial PV energy coverage for {title} {year:04d}: "
                        f"covered_days={int(values['covered_days'])}, expected_days={int(year_expected_days)}, ratio={ratio:.3f}"
                    )
            energy_wh = values["energy_wh"]
            cost_eur = values["covered_cost_eur"]
            peak_labels = {
                "title": title,
                "local_year": f"{year:04d}",
            }
            installed_watt_peak = average_installed_watt_peak_for_calendar_year(title_assets, year)
            if energy_wh > 1000 and ratio is not None and ratio >= partial_warning_threshold:
                append_sample(series, "evcc_pv_energy_by_title_yearly_wh", peak_labels, timestamp_for_year(year, tz), energy_wh)
                energy_coverage_labels = {
                    "title": title,
                    "coverage": coverage_label(ratio),
                    "local_year": f"{year:04d}",
                }
                append_sample(series, "evcc_pv_energy_by_title_yearly_with_coverage_wh", energy_coverage_labels, timestamp_for_year(year, tz), energy_wh)
            if energy_wh > 1000 and installed_watt_peak > 0:
                yield_labels = {
                    "title": title,
                    "coverage": coverage_label(ratio),
                    "local_year": f"{year:04d}",
                }
                append_sample(series, "evcc_pv_specific_yield_yearly_with_coverage_kwh_per_kwp", yield_labels, timestamp_for_year(year, tz), energy_wh / installed_watt_peak)
                if ratio >= partial_warning_threshold:
                    append_sample(series, "evcc_pv_installed_watt_peak_yearly", peak_labels, timestamp_for_year(year, tz), installed_watt_peak)
                    append_sample(series, "evcc_pv_specific_yield_yearly_kwh_per_kwp", peak_labels, timestamp_for_year(year, tz), energy_wh / installed_watt_peak)

            if energy_wh <= 1000 or cost_eur <= 0 or not should_write_lcoe(ratio, min_lcoe_coverage_ratio):
                continue

            labels = {
                "title": title,
                "coverage": coverage_label(ratio),
                "local_year": f"{year:04d}",
            }
            append_sample(series, "evcc_pv_lcoe_yearly_ct_per_kwh", labels, timestamp_for_year(year, tz), cost_eur / (energy_wh / 1000.0) * 100.0)

        for day in iter_days(start_day, end_day):
            window_start = day - dt.timedelta(days=6)
            window_energy_wh = 0.0
            window_cost_eur = 0.0
            window_expected_days = 0.0
            window_covered_days = 0.0
            current = window_start
            while current <= day:
                values = daily_title_values.get(current)
                if values:
                    window_energy_wh += values["energy_wh"]
                    window_cost_eur += values["covered_cost_eur"]
                    window_expected_days += values["expected_days"]
                    window_covered_days += values["covered_days"]
                current += dt.timedelta(days=1)
            ratio = coverage_ratio(window_covered_days, window_expected_days)
            if window_energy_wh > 1000:
                labels = {
                    "title": title,
                    "local_year": f"{day.year:04d}",
                }
                window_watt_peak = average_installed_watt_peak_for_range(title_assets, window_start, day + dt.timedelta(days=1))
                if window_watt_peak > 0:
                    append_sample(series, "evcc_pv_specific_yield_rolling_7d_kwh_per_kwp", labels, timestamp_for_day(day, tz), window_energy_wh / window_watt_peak)
                if window_cost_eur > 0 and should_write_lcoe(ratio, min_lcoe_coverage_ratio):
                    append_sample(series, "evcc_pv_lcoe_rolling_7d_ct_per_kwh", labels, timestamp_for_day(day, tz), window_cost_eur / (window_energy_wh / 1000.0) * 100.0)

        for (asset_id, (year, month)), cost_eur in sorted(monthly_asset_costs.items()):
            labels = {
                "asset_id": asset_id,
                "title": title,
                "local_year": f"{year:04d}",
            }
            append_sample(series, "evcc_pv_investment_cost_monthly_eur", labels, timestamp_for_month(year, month, tz), cost_eur)

        for (year, month), values in sorted(monthly_title_totals.items()):
            energy_wh = values["energy_wh"]
            cost_eur = values["covered_cost_eur"]
            if energy_wh <= 0:
                continue
            labels = {
                "title": title,
                "local_year": f"{year:04d}",
            }
            ts = timestamp_for_month(year, month, tz)
            append_sample(series, "evcc_pv_lcoe_energy_monthly_wh", labels, ts, energy_wh)
            ratio = coverage_ratio(values["covered_days"], values["expected_days"])
            if cost_eur > 0 and should_write_lcoe(ratio, min_lcoe_coverage_ratio):
                append_sample(series, "evcc_pv_lcoe_cost_monthly_eur", labels, ts, cost_eur)
                append_sample(series, "evcc_pv_lcoe_monthly_ct_per_kwh", labels, ts, cost_eur / (energy_wh / 1000.0) * 100.0)


        title_ratio = coverage_ratio(title_covered_days, title_calendar_expected_days or title_expected_days)
        title_summaries.append({
            "title": title,
            "investment_rows": len(title_assets),
            "energy_kwh": title_energy_wh / 1000.0,
            "active_cost_eur": title_active_cost_eur,
            "covered_cost_eur": title_covered_cost_eur,
            "cost_eur": title_covered_cost_eur,
            "coverage_ratio": title_ratio,
            "covered_days": int(title_covered_days),
            "expected_days": int(title_calendar_expected_days or title_expected_days),
            "partial": (title_ratio is not None and title_ratio < partial_warning_threshold),
            "lcoe_ct_per_kwh": (title_covered_cost_eur / (title_energy_wh / 1000.0) * 100.0) if title_energy_wh > 0 and should_write_lcoe(title_ratio, min_lcoe_coverage_ratio) else None,
            "energy_merge": energy_merge_stats,
        })
        for asset in title_assets:
            asset_active_cost_eur = asset_active_costs[asset["asset_id"]]
            asset_covered_cost_eur = asset_covered_costs[asset["asset_id"]]
            asset_summaries.append({
                "asset_id": asset["asset_id"],
                "title": title,
                "energy_kwh": title_energy_wh / 1000.0,
                "active_cost_eur": asset_active_cost_eur,
                "covered_cost_eur": asset_covered_cost_eur,
                "cost_eur": asset_covered_cost_eur,
                "coverage_ratio": title_ratio,
                "lcoe_ct_per_kwh": (asset_covered_cost_eur / (title_energy_wh / 1000.0) * 100.0) if title_energy_wh > 0 and should_write_lcoe(title_ratio, min_lcoe_coverage_ratio) else None,
            })

    if write_pv_energy_rollup:
        source_label = pv_energy_rollup_source_label(energy_source)
        for day, values in sorted(totals.items()):
            energy_wh = values.get("energy_wh", 0.0)
            if energy_wh <= 0:
                continue
            labels = {
                "source": source_label,
                "local_year": f"{day.year:04d}",
                "local_month": f"{day.month:02d}",
            }
            append_sample(series, "evcc_pv_energy_daily_wh", labels, timestamp_for_day(day, tz), energy_wh)

    yearly_totals: dict[int, dict[str, float]] = {}
    for (_year, _month), values in monthly_totals.items():
        yearly_totals.setdefault(_year, {"energy_wh": 0.0, "cost_eur": 0.0})
        yearly_totals[_year]["energy_wh"] += values["energy_wh"]
        yearly_totals[_year]["cost_eur"] += values["cost_eur"]


    for year, values in sorted(yearly_totals.items()):
        energy_wh = values["energy_wh"]
        cost_eur = values["cost_eur"]
        if energy_wh <= 1000 or cost_eur <= 0:
            continue
        labels = {
            "scope": "pv",
            "local_year": f"{year:04d}",
        }
        append_sample(series, "evcc_pv_effective_lcoe_yearly_ct_per_kwh", labels, timestamp_for_year(year, tz), cost_eur / (energy_wh / 1000.0) * 100.0)
    for (year, month), values in sorted(monthly_totals.items()):
        energy_wh = values["energy_wh"]
        cost_eur = values["cost_eur"]
        if energy_wh <= 0:
            continue
        labels = {
            "scope": "pv",
            "local_year": f"{year:04d}",
        }
        append_sample(series, "evcc_pv_effective_lcoe_monthly_ct_per_kwh", labels, timestamp_for_month(year, month, tz), cost_eur / (energy_wh / 1000.0) * 100.0)

    for day, values in sorted(totals.items()):
        energy_wh = values["energy_wh"]
        cost_eur = values["cost_eur"]
        if energy_wh <= 0:
            continue
        labels = {
            "scope": "pv",
            "local_year": f"{day.year:04d}",
            "local_month": f"{day.month:02d}",
        }
        append_sample(series, "evcc_pv_effective_lcoe_daily_ct_per_kwh", labels, timestamp_for_day(day, tz), cost_eur / (energy_wh / 1000.0) * 100.0)

    summary = {
        "assets": asset_summaries,
        "pv_sources": title_summaries,
        "coverage": coverage_summaries,
        "shared_allocations": shared_allocation_summaries,
        "series": len(series),
        "samples": sum(len(v) for v in series.values()),
        "start": start_day.isoformat(),
        "end_exclusive": end_day.isoformat(),
        "write_pv_energy_rollup": write_pv_energy_rollup,
        "min_lcoe_coverage_ratio": min_lcoe_coverage_ratio,
        "partial_warning_threshold": partial_warning_threshold,
    }
    return series, summary

def import_series(base_url: str, series: dict[tuple[str, tuple[tuple[str, str], ...]], list[tuple[int, float]]], batch_size: int) -> int:
    batch: list[str] = []
    batches = 0
    for (metric, label_items), samples in series.items():
        labels = dict(label_items)
        samples = sorted(samples)
        batch.append(line(metric, labels, samples))
        if len(batch) >= batch_size:
            post_bytes(f"{base_url.rstrip('/')}/api/v1/import", ("\n".join(batch) + "\n").encode("utf-8"))
            batches += 1
            batch = []
    if batch:
        post_bytes(f"{base_url.rstrip('/')}/api/v1/import", ("\n".join(batch) + "\n").encode("utf-8"))
        batches += 1
    return batches


def main() -> int:
    parser = argparse.ArgumentParser(description="Import EVCC PV investment cost rollups into VictoriaMetrics")
    parser.add_argument("--vm-base-url", default="http://127.0.0.1:8428")
    parser.add_argument("--investment-file", default="data/private/investments.xlsx")
    parser.add_argument("--start", help="first local day, YYYY-MM-DD")
    parser.add_argument("--end", help="exclusive local end day, YYYY-MM-DD")
    parser.add_argument("--timezone", default="Europe/Berlin")
    parser.add_argument("--peak-power-limit", type=float, default=30000.0)
    parser.add_argument("--sample-interval", default="30s")
    parser.add_argument("--energy-source", choices=["pv-power", "daily-metric", "combined"], default="pv-power", help="pv-power integrates EVCC pvPower_value; daily-metric reads preimported daily Wh metrics; combined merges both")
    parser.add_argument("--pv-energy-metric", default="evcc_pv_energy_by_title_daily_wh", help="per-title daily Wh metric to use with --energy-source=daily-metric or --energy-source=combined")
    parser.add_argument("--skip-titles-without-energy", action="store_true", help="skip investment titles with no energy samples in the selected energy source")
    parser.add_argument("--combined-energy-conflict", choices=["prefer-evcc", "prefer-daily-metric", "sum", "error"], default="prefer-evcc", help="how --energy-source=combined handles days where EVCC pvPower and the daily metric both have energy")
    parser.add_argument("--write-pv-energy-rollup", action="store_true", help="also write evcc_pv_energy_daily_wh from the selected daily PV energy source for standard PV dashboard panels")
    parser.add_argument("--min-lcoe-coverage-ratio", type=float, default=0.0, help="minimum covered/expected cost-day ratio required before writing monthly/yearly/rolling LCOE values")
    parser.add_argument("--partial-warning-threshold", type=float, default=0.95, help="coverage ratio below which summaries and warnings mark a period as partial")
    parser.add_argument("--batch-size", type=int, default=200)
    parser.add_argument("--replace", action="store_true", help="delete generated evcc_pv_* investment metrics before writing")
    parser.add_argument("--write", action="store_true", help="write generated metrics to VictoriaMetrics")
    args = parser.parse_args()

    started = time.perf_counter()
    log(f"import-investment-costs.py v{SCRIPT_VERSION} (last modified {SCRIPT_LAST_MODIFIED})")
    path = Path(args.investment_file)
    if not path.exists():
        raise SystemExit(f"Investment file not found: {path}")

    tz = ZoneInfo(args.timezone)
    assets = read_investments(path)
    if not assets:
        raise SystemExit(f"No investment rows found in {path}")
    pv_assets = [a for a in assets if a["asset_type"] == "pv" and a["include"]]
    pv_cost_assets = [a for a in assets if a["asset_type"] in {"pv", "pv_shared"} and a["include"]]
    if not pv_assets:
        raise SystemExit("No included PV energy assets found. Set asset_type=pv and include_in_effective_price=yes.")

    default_start = min(a["commissioning_date"] for a in pv_assets)
    start_day = dt.date.fromisoformat(args.start) if args.start else default_start
    end_day = dt.date.fromisoformat(args.end) if args.end else dt.datetime.now(tz).date() + dt.timedelta(days=1)
    if end_day <= start_day:
        raise SystemExit("--end must be after --start")

    log(f"PV energy assets: {len(pv_assets)}")
    log(f"PV cost rows: {len(pv_cost_assets)}")
    log(f"Range: {start_day} .. {end_day} (exclusive), timezone={args.timezone}")
    log(f"Energy source: {args.energy_source}" + (f" ({args.pv_energy_metric})" if args.energy_source in {"daily-metric", "combined"} else ""))
    series, summary = build_rollups(args.vm_base_url, assets, start_day, end_day, tz, args.peak_power_limit, args.sample_interval, args.energy_source, args.pv_energy_metric, args.skip_titles_without_energy, args.write_pv_energy_rollup, args.combined_energy_conflict, args.min_lcoe_coverage_ratio, args.partial_warning_threshold)
    log(json.dumps(summary, indent=2, ensure_ascii=False))

    if args.write:
        if args.replace:
            for metric in GENERATED_METRICS + LEGACY_GENERATED_METRICS:
                log(f"Delete existing metric: {metric}")
                delete_metric(args.vm_base_url, metric)
            if args.write_pv_energy_rollup:
                selector = f'evcc_pv_energy_daily_wh{{source="{pv_energy_rollup_source_label(args.energy_source)}"}}'
                log(f"Delete existing helper PV energy rollup: {selector}")
                delete_series(args.vm_base_url, selector)
        batches = import_series(args.vm_base_url, series, args.batch_size)
        log(f"Imported batches: {batches}")
    else:
        log("Dry run only. Re-run with --write to import metrics.")

    log(f"Finished in {time.perf_counter() - started:.1f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

